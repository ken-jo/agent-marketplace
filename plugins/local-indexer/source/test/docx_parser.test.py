"""Tests for DOCX multi-page parsing in LocalIndexer._parse_docx.

Verifies:
1. Multi-page DOCX with hard page breaks → content bucketed into correct pages.
2. Multi-page DOCX with page_break_before format → content bucketed correctly.
3. Single-page DOCX with no breaks → returns single page, logs the condition.
4. All text is preserved (no content loss) across all pages.
5. Tables are included in page content.
"""

from __future__ import annotations

import io
import logging
import sys
from pathlib import Path
from typing import List, Tuple
from unittest.mock import patch

import pytest

_WORKER_DIR = Path(__file__).resolve().parent.parent / "worker"
sys.path.insert(0, str(_WORKER_DIR))

python_docx = pytest.importorskip("docx", reason="python-docx not installed")

from docx import Document  # noqa: E402 — after importorskip
from local_indexer import LocalIndexer  # noqa: E402


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_indexer(tmp_path: Path) -> LocalIndexer:
    """Minimal LocalIndexer with no embedding/vector backends."""
    return LocalIndexer(
        workspace=str(tmp_path / "workspace"),
        sqlite_path=str(tmp_path / "fts5.sqlite"),
        lance_path=str(tmp_path / "vectors.lance"),
        openai_api_key=None,
        embedding_client=None,
        vector_store=None,
    )


def _build_docx_hard_breaks(page_texts: List[str]) -> bytes:
    """Create a DOCX in-memory with hard page breaks between each text block.

    The first text block goes on page 1; subsequent blocks are preceded by
    ``doc.add_page_break()`` which inserts a ``<w:br w:type="page"/>`` run.
    Returns the raw DOCX bytes so the caller can write it to a temp file.
    """
    doc = Document()
    for i, text in enumerate(page_texts):
        if i > 0:
            doc.add_page_break()
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_page_break_before(page_texts: List[str]) -> bytes:
    """Create a DOCX with ``page_break_before=True`` on each paragraph after the first."""
    doc = Document()
    for i, text in enumerate(page_texts):
        para = doc.add_paragraph(text)
        if i > 0:
            para.paragraph_format.page_break_before = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_no_breaks(text: str) -> bytes:
    """Create a single-paragraph DOCX with no page breaks."""
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_with_table(page1_text: str, table_rows: List[Tuple[str, str]], page2_text: str) -> bytes:
    """Create a DOCX with a table on page 1, then a hard break, then page 2."""
    doc = Document()
    doc.add_paragraph(page1_text)
    tbl = doc.add_table(rows=len(table_rows), cols=2)
    for i, (c1, c2) in enumerate(table_rows):
        tbl.rows[i].cells[0].text = c1
        tbl.rows[i].cells[1].text = c2
    doc.add_page_break()
    doc.add_paragraph(page2_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestDocxHardPageBreaks:
    """Hard breaks inserted via add_page_break() → w:br type='page'."""

    def test_two_pages_bucketed_correctly(self, tmp_path: Path) -> None:
        """Content on page 1 and page 2 must be in separate (page, text) tuples."""
        docx_bytes = _build_docx_hard_breaks(["First page content", "Second page content"])
        docx_path = tmp_path / "two_pages.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, page_count = indexer._parse_docx(str(docx_path))

        page_nums = [p for p, _ in pages]
        page_texts = [t for _, t in pages]

        assert 1 in page_nums, "page 1 must exist"
        assert 2 in page_nums, "page 2 must exist"
        assert len(pages) == 2, "exactly two pages expected"

        p1_text = next(t for p, t in pages if p == 1)
        p2_text = next(t for p, t in pages if p == 2)
        assert "First page content" in p1_text
        assert "Second page content" in p2_text

    def test_three_pages_bucketed_correctly(self, tmp_path: Path) -> None:
        docx_bytes = _build_docx_hard_breaks(["Alpha", "Beta", "Gamma"])
        docx_path = tmp_path / "three_pages.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, page_count = indexer._parse_docx(str(docx_path))

        assert len(pages) == 3
        assert pages[0][0] == 1
        assert pages[1][0] == 2
        assert pages[2][0] == 3
        assert "Alpha" in pages[0][1]
        assert "Beta" in pages[1][1]
        assert "Gamma" in pages[2][1]

    def test_no_content_loss(self, tmp_path: Path) -> None:
        """All text from all pages must be present in the aggregate output."""
        expected_texts = ["First page content", "Second page content", "Third page content"]
        docx_bytes = _build_docx_hard_breaks(expected_texts)
        docx_path = tmp_path / "no_loss.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, _ = indexer._parse_docx(str(docx_path))

        all_text = " ".join(t for _, t in pages)
        for expected in expected_texts:
            assert expected in all_text, f"text must not be lost: {expected!r}"

    def test_page_count_returned(self, tmp_path: Path) -> None:
        docx_bytes = _build_docx_hard_breaks(["P1", "P2", "P3"])
        docx_path = tmp_path / "pc.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, page_count = indexer._parse_docx(str(docx_path))

        assert page_count == 3, "page_count must equal max page number"


class TestDocxPageBreakBeforeFormat:
    """Hard breaks via paragraph.paragraph_format.page_break_before = True."""

    def test_two_pages_bucketed_correctly(self, tmp_path: Path) -> None:
        docx_bytes = _build_docx_page_break_before(["Intro paragraph", "Body paragraph"])
        docx_path = tmp_path / "pbf.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, _ = indexer._parse_docx(str(docx_path))

        assert len(pages) == 2
        p1 = next(t for p, t in pages if p == 1)
        p2 = next(t for p, t in pages if p == 2)
        assert "Intro paragraph" in p1
        assert "Body paragraph" in p2

    def test_no_content_loss(self, tmp_path: Path) -> None:
        texts = ["Section A", "Section B", "Section C"]
        docx_bytes = _build_docx_page_break_before(texts)
        docx_path = tmp_path / "pbf_no_loss.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, _ = indexer._parse_docx(str(docx_path))

        all_text = " ".join(t for _, t in pages)
        for expected in texts:
            assert expected in all_text


class TestDocxNoPageBreaks:
    """DOCX with no page breaks → single page, condition logged."""

    def test_single_page_fallback(self, tmp_path: Path) -> None:
        docx_bytes = _build_docx_no_breaks("Only page content here")
        docx_path = tmp_path / "single.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, page_count = indexer._parse_docx(str(docx_path))

        assert len(pages) == 1, "single-page DOCX must yield exactly one page"
        assert pages[0][0] == 1, "single page must be numbered 1"
        assert "Only page content here" in pages[0][1]

    def test_no_breaks_is_logged(self, tmp_path: Path) -> None:
        """The no-breaks condition must be logged (not silent)."""
        docx_bytes = _build_docx_no_breaks("Text on the only page")
        docx_path = tmp_path / "logged.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)

        log_calls: list = []

        original_info = None
        import local_indexer as _li_mod

        _original_info = _li_mod._info  # type: ignore[attr-defined]

        def _capture_info(logger: object, event: str, **kw: object) -> None:
            log_calls.append(event)
            _original_info(logger, event, **kw)

        with patch.object(_li_mod, "_info", side_effect=_capture_info):
            indexer._parse_docx(str(docx_path))

        assert "indexer.docx.no_page_breaks" in log_calls, (
            "no-page-breaks condition must emit 'indexer.docx.no_page_breaks' log event"
        )

    def test_multi_page_has_no_spurious_log(self, tmp_path: Path) -> None:
        """A DOCX with detected breaks must NOT emit the no-breaks log event."""
        docx_bytes = _build_docx_hard_breaks(["Page 1", "Page 2"])
        docx_path = tmp_path / "has_breaks.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        import local_indexer as _li_mod

        log_calls: list = []
        _original_info = _li_mod._info  # type: ignore[attr-defined]

        def _capture_info(logger: object, event: str, **kw: object) -> None:
            log_calls.append(event)
            _original_info(logger, event, **kw)

        with patch.object(_li_mod, "_info", side_effect=_capture_info):
            indexer._parse_docx(str(docx_path))

        assert "indexer.docx.no_page_breaks" not in log_calls


class TestDocxWithTable:
    """Table content must be included in page output."""

    def test_table_content_on_correct_page(self, tmp_path: Path) -> None:
        table_rows = [("Name", "Value"), ("Row1", "Data1")]
        docx_bytes = _build_docx_with_table(
            "Before table", table_rows, "After page break"
        )
        docx_path = tmp_path / "table.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, _ = indexer._parse_docx(str(docx_path))

        assert len(pages) >= 2, "table + break must yield at least 2 pages"

        p1_text = next(t for p, t in pages if p == 1)
        p2_text = next(t for p, t in pages if p == 2)

        assert "Before table" in p1_text
        assert "Name" in p1_text or "Row1" in p1_text, "table cell content must appear on page 1"
        assert "After page break" in p2_text

    def test_no_content_loss_with_table(self, tmp_path: Path) -> None:
        table_rows = [("Col A", "Col B"), ("r1c1", "r1c2")]
        docx_bytes = _build_docx_with_table("Intro", table_rows, "Conclusion")
        docx_path = tmp_path / "table_loss.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, _ = indexer._parse_docx(str(docx_path))

        all_text = " ".join(t for _, t in pages)
        assert "Intro" in all_text
        assert "Conclusion" in all_text
        # At least some table cell content must survive
        assert "Col A" in all_text or "r1c1" in all_text


class TestDocxIntegration:
    """End-to-end: _parse_docx on a multi-page DOCX returns correct page structure."""

    def test_parse_docx_returns_multiple_pages(self, tmp_path: Path) -> None:
        """_parse_docx must return multiple (page, text) pairs — not all page=1."""
        docx_bytes = _build_docx_hard_breaks([
            "Alpha content on the first page",
            "Beta content on the second page",
            "Gamma content on the third page",
        ])
        docx_path = tmp_path / "integration.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = _make_indexer(tmp_path)
        pages, page_count = indexer._parse_docx(str(docx_path))

        page_nums = {p for p, _ in pages}
        assert page_nums != {1}, (
            "_parse_docx must not collapse all content to page=1. "
            f"Got page numbers: {sorted(page_nums)}"
        )
        assert 1 in page_nums, "page 1 must be present"
        assert 2 in page_nums, "page 2 must be present"
        assert 3 in page_nums, "page 3 must be present"
        assert page_count == 3

        # Verify content is on the correct pages
        page_map = {p: t for p, t in pages}
        assert "Alpha" in page_map[1]
        assert "Beta" in page_map[2]
        assert "Gamma" in page_map[3]

    def test_index_file_multi_page_docx_succeeds(self, tmp_path: Path) -> None:
        """index_file on a multi-page DOCX must succeed with page_count > 1."""
        from typing import Any, Dict, List as TList

        class _FakeEmbed:
            dimensions = 1536
            def embed_corpus(self, chunks: TList[str]) -> TList[TList[float]]:
                return [[0.01] * 1536 for _ in chunks]
            def embed_one(self, text: str) -> TList[float]:
                return [0.01] * 1536

        class _FakeStore:
            def __init__(self) -> None:
                self.rows: TList[Dict[str, Any]] = []
            def upsert(self, rows: TList[Dict[str, Any]]) -> None:
                self.rows.extend(rows)
            def search(self, vector: TList[float], top_k: int = 20) -> TList[Dict[str, Any]]:
                return []

        class _FakeTokenizer:
            def tokenize_for_index(self, text: str) -> str:
                return " ".join(text.split())
            def split_sents(self, text: str) -> List[str]:
                return [text] if text.strip() else []

        docx_bytes = _build_docx_hard_breaks([
            "First page text content here",
            "Second page text content here",
        ])
        docx_path = tmp_path / "e2e.docx"
        docx_path.write_bytes(docx_bytes)

        indexer = LocalIndexer(
            workspace=str(tmp_path / "workspace"),
            sqlite_path=str(tmp_path / "fts5.sqlite"),
            lance_path=str(tmp_path / "vectors.lance"),
            openai_api_key="sk-test",
            embedding_client=_FakeEmbed(),
            vector_store=_FakeStore(),
            tokenizer=_FakeTokenizer(),
        )

        result = indexer.index_file(str(docx_path), mode="docx")
        assert result.status == "indexed"
        assert result.chunk_count > 0
        assert result.page_count == 2, (
            f"page_count must be 2 for a 2-page DOCX, got {result.page_count}"
        )
